import json,re,hashlib,math
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parents[1];up=root.parent/'upload'
def norm(s):return re.sub(r'[\s\u200c-\u200f\u202a-\u202e]','',str(s).replace('ي','ی').replace('ك','ک'))
geo=json.loads((root/'public/data/iran-provinces.geojson').read_text());provinces={norm(f['properties']['name']):f['properties']['ISO3166-2'] for f in geo['features']}
provinces[norm('کهکیلویه و بویراحمد')]='IR-18'
speed=[];port=[];checks=[];origins=[]
for filename in ['Speed _ Fcp & Unsp _ 1405.06.xlsx','سرعت 06-1405.xlsx','Port _ Fcp & Unsp _ 1405.06.xlsx','پورت 06-1405.xlsx']:
 p=up/filename;origins.append(dict(file=filename,sha256=hashlib.sha256(p.read_bytes()).hexdigest()));w=openpyxl.load_workbook(p,read_only=True,data_only=True)
 is_speed='peed' in filename or 'سرعت' in filename
 for s in w:
  a=list(s.values)
  if is_speed:
   hdr=next((i for i,r in enumerate(a[:6]) if len(r)>1 and 'نام شرکت' in str(r[1])),None)
   assert hdr is not None,(filename,s.title)
   data=[];total=None
   for i,r in enumerate(a[hdr+2:],hdr+3):
    if isinstance(r[0],(int,float)) and isinstance(r[1],str):data.append(dict(company=r[1].strip(),row=i,values=[v if isinstance(v,(int,float)) else None for v in r[2:26]]))
    if any(norm(v) in ['جمعکل','مجموع'] for v in r[:2] if v is not None):total=dict(company='جمع گزارش‌شده',row=i,values=[v if isinstance(v,(int,float)) else None for v in r[2:26]])
   bands=[dict(label=str(a[hdr][c]).strip(),source_column=c+1,total_index=c,person_index=c-2,legal_index=c-1) for c in range(2,23,3)]
   # values array starts in column C; band total is E,H,...,W => index 2,5,...,20.
   issue=[]
   for row in data+([total] if total else []):
    for j,v in enumerate(row['values']):
     if isinstance(v,(int,float)) and not float(v).is_integer():issue.append(dict(row=row['row'],cell=f'{openpyxl.utils.get_column_letter(j+3)}{row["row"]}',value=v,reason='FRACTIONAL_SUBSCRIBER_COUNT'))
   deltas=[]
   if total:
    for j,t in enumerate(total['values']):
     vals=[r['values'][j] for r in data]
     if t is not None and all(v is not None for v in vals):
      delta=sum(vals)-t
      if abs(delta)>1e-6:deltas.append(dict(column=j+3,delta=delta))
   speed.append(dict(id=f'speed-{len(speed)}',file=filename,sheet=s.title,header_row=hdr+1,period='شهریور ۱۴۰۵؛ بر اساس نام فایل، مگر ذکر صریح در عنوان',bands=bands,rows=data,total=total,issues=issue,reconciliation_differences=deltas,scope='شرکت‌های همین فایل؛ باند سرعت سرویس، نه سرعت اندازه‌گیری‌شده'))
  else:
   if s.title=='تجمیع':checks.append(dict(file=filename,sheet=s.title,status='RAW_ONLY_OVERLAPPING_AGGREGATION'));continue
   fcp=filename.startswith('Port');company_header=2 if fcp else 1;metric_header=3 if fcp else 2;start=4 if fcp else 3
   cols=[];provider=''
   for c in range(1,s.max_column):
    if c<len(a[company_header]) and a[company_header][c] is not None:provider=str(a[company_header][c]).strip()
    if not fcp and c<4:provider='مجموع گزارش‌شده فایل'
    metric=(a[metric_header][c] if c<len(a[metric_header]) else None) if fcp or c>=4 else a[0][c]
    if isinstance(metric,str) and norm(metric) in ['کلمنصوبه','کلدایری','دایری','حقیقی','حقوقی','خانگی','تجاری']:
     cols.append(dict(column=c+1,letter=openpyxl.utils.get_column_letter(c+1),provider=provider,metric=metric.strip(),aggregate=('مجموع' in provider),active=('دایری' in norm(metric))))
   rows=[];total=None;unmapped=[]
   for i,r in enumerate(a[start:],start+1):
    name=str(r[0] or '').strip()
    if norm(name) in ['کل','جمعکل','مجموع']:
     total=dict(row=i,values=[r[c['column']-1] if isinstance(r[c['column']-1],(int,float)) else None for c in cols]);break
    if not name or 'توضیحات' in name:continue
    iso=provinces.get(norm(name))
    if not iso:unmapped.append(name);continue
    rows.append(dict(province=name,iso=iso,row=i,values=[r[c['column']-1] if isinstance(r[c['column']-1],(int,float)) else None for c in cols]))
   deltas=[]
   if total:
    for j,t in enumerate(total['values']):
     vals=[r['values'][j] for r in rows]
     if t is not None and all(v is not None for v in vals) and abs(sum(vals)-t)>1e-6:deltas.append(dict(column=cols[j]['letter'],delta=sum(vals)-t))
   assert len({r['iso'] for r in rows})==len(rows),(filename,s.title,'duplicate provinces')
   port.append(dict(id=f'port-{len(port)}',file=filename,sheet=s.title,period='شهریور ۱۴۰۵؛ فایل‌های ارسالی',columns=cols,rows=rows,total=total,unmapped=unmapped,reconciliation_differences=deltas,scope='اعلام همین فایل؛ مجموع‌ها و شرکت‌ها هم‌پوشانی دارند'))
result=dict(speed=speed,ports=port,raw_only=checks,sources=origins,normalization='Arabic/Persian character and whitespace normalization; province alias کهکیلویه mapped explicitly to IR-18; no geography inferred from IP',validation=dict(speed_tables=len(speed),port_tables=len(port),speed_rows=sum(len(t['rows']) for t in speed),port_province_rows=sum(len(t['rows']) for t in port),fractional_cells=sum(len(t['issues']) for t in speed),unmapped_provinces=sum(len(t['unmapped']) for t in port)))
assert result['validation']['speed_tables']==12
assert result['validation']['port_tables']==14
assert result['validation']['unmapped_provinces']==0,result['validation']
(root/'src/command-access-normalized.json').write_text(json.dumps(result,ensure_ascii=False,separators=(',',':')))
(root/'docs/audits/access-normalization-summary.json').write_text(json.dumps(dict(validation=result['validation'],raw_only=checks,sources=origins,port_checks=[dict(id=t['id'],file=t['file'],sheet=t['sheet'],provinces=len(t['rows']),deltas=len(t['reconciliation_differences'])) for t in port],speed_checks=[dict(id=t['id'],file=t['file'],sheet=t['sheet'],companies=len(t['rows']),issues=len(t['issues']),deltas=len(t['reconciliation_differences'])) for t in speed]),ensure_ascii=False,indent=2))
print(result['validation'])
